import io
import unittest
import zipfile

import core
from docx import Document


class BulkGenerationTests(unittest.TestCase):
    def test_classifies_customer_from_first_fiscal_code_character(self):
        self.assertEqual(core.classify_customer_type("RSSMRA80A01H501U"), "fisica")
        self.assertEqual(core.classify_customer_type("12345678901"), "giuridica")
        self.assertIsNone(core.classify_customer_type(""))
        self.assertIsNone(core.classify_customer_type("-1234567890"))

    def test_splits_additional_legal_representative_and_owner_placeholders(self):
        self.assertEqual(core.split_role("COGNOME_LR2"), ("COGNOME", "LR2"))
        self.assertEqual(core.split_role("PERCENTUALE_TE3"), ("PERCENTUALE", "TE3"))
        self.assertEqual(core.role_label("LR2"), "Legale rappresentante 2")
        self.assertEqual(core.role_label("TE3"), "Titolare effettivo 3")

    def test_formats_percentages_as_italian_values_with_two_decimals(self):
        self.assertEqual(core.format_percentage_value("10"), "10,00%")
        self.assertEqual(core.format_percentage_value("10,5"), "10,50%")
        self.assertEqual(core.format_percentage_value("10.50%"), "10,50%")
        with self.assertRaisesRegex(ValueError, "compresa tra 0 e 100"):
            core.format_percentage_value("100,01")

    def test_finds_and_fills_fields_inside_archived_premise(self):
        premise = "La società [SOCIETA] con sede in [INDIRIZZO], rappresentata da [SOCIETA]."
        self.assertEqual(
            core.find_premise_fields(premise),
            ["SOCIETA", "INDIRIZZO"],
        )
        filled = core.fill_premise_fields(
            premise,
            {"SOCIETA": "Alfa S.r.l.", "INDIRIZZO": "Via Roma 1"},
        )
        self.assertEqual(
            filled,
            "La società Alfa S.r.l. con sede in Via Roma 1, "
            "rappresentata da Alfa S.r.l..",
        )

    def test_formats_numeric_premise_fields_as_amounts_and_words(self):
        filled = core.fill_premise_fields(
            "Il compenso è [IMPORTO] ([IMPORTO]), la quota è "
            "[PERCENTUALE] e il motivo è [CAUSALE].",
            {
                "IMPORTO": "1000",
                "PERCENTUALE": "10",
                "CAUSALE": "consulenza",
            },
        )

        self.assertEqual(
            filled,
            "Il compenso è 1.000,00 (mille/00) "
            "(1.000,00 (mille/00)), la quota è "
            "10,00% (dieci per cento) e il motivo è consulenza.",
        )
        self.assertEqual(
            core.format_premise_field_value("IMPORTO", "1000,5"),
            "1.000,50 (mille/50)",
        )
        with self.assertRaisesRegex(ValueError, "importo numerico"):
            core.format_premise_field_value("IMPORTO", "1000,555")
        self.assertEqual(
            core.format_premise_field_value("PERCENTUALE", "10,5"),
            "10,50% (dieci virgola cinquanta per cento)",
        )
        self.assertEqual(
            core.format_premise_field_value("PERCENTUALE", "10,05"),
            "10,05% (dieci virgola zero cinque per cento)",
        )
        with self.assertRaisesRegex(ValueError, "compresa tra 0 e 100"):
            core.format_premise_field_value("PERCENTUALE", "100,01")

    def test_filled_premise_fields_are_rendered_in_document(self):
        document = Document()
        document.add_paragraph("Premessa: {{PREMESSA}}")
        source = io.BytesIO()
        document.save(source)
        premise = core.fill_premise_fields(
            "La società [SOCIETA] è rappresentata da [LEGALE RAPPRESENTANTE].",
            {
                "SOCIETA": "Alfa S.r.l.",
                "LEGALE RAPPRESENTANTE": "Mario Rossi",
            },
        )

        rendered = core.render(
            source.getvalue(),
            [{"name": "PREMESSA", "tokens": ["{{PREMESSA}}"]}],
            {"PREMESSA": premise},
        )

        output = Document(io.BytesIO(rendered))
        self.assertEqual(
            output.paragraphs[0].text,
            "Premessa: La società Alfa S.r.l. è rappresentata da Mario Rossi.",
        )

    def test_creates_zip_with_individual_documents(self):
        content = core.create_document_archive(
            [("Mandato Mario Rossi.docx", b"persona fisica"),
             ("Mandato Societa Uno.docx", b"persona giuridica")]
        )

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            self.assertEqual(
                archive.namelist(),
                ["Mandato Mario Rossi.docx", "Mandato Societa Uno.docx"],
            )
            self.assertEqual(
                archive.read("Mandato Mario Rossi.docx"), b"persona fisica"
            )
            self.assertEqual(
                archive.read("Mandato Societa Uno.docx"), b"persona giuridica"
            )


if __name__ == "__main__":
    unittest.main()
