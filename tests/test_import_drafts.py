import tempfile
import unittest
from pathlib import Path

import core


class ChecklistTests(unittest.TestCase):
    def _checklist(self):
        import io
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Lista interna"
        for column, header in enumerate(
            ["# Mandato", "Cliente", "% capitale sociale", "Commissione Apertura", "Totale", "Stato"], 2
        ):
            sheet.cell(3, column, header)
        sheet.append([None, "x", "mario rossi", 0.375, 1500, "#REF!", "pronto"])
        sheet.cell(sheet.max_row, 4).number_format = "0.00%"
        sheet.append([None, "y", "TOTALE", 1, 3000])
        out = io.BytesIO()
        workbook.save(out)
        return core.read_checklist(out.getvalue(), "lista.xlsx")

    def test_reads_header_offset_percentages_errors(self):
        checklist = self._checklist()
        headers = checklist["headers"]
        self.assertEqual(core.find_checklist_column(headers, core.CHECKLIST_NAME_HEADERS), 2)
        first, total = checklist["rows"]
        self.assertEqual(first["cells"][2], "mario rossi")
        self.assertEqual(first["cells"][3], "37.5%")
        self.assertEqual(first["cells"][5], "")
        self.assertTrue(core.is_checklist_total_row(total["cells"][2]))

    def test_name_matching(self):
        candidates = [
            ("a", "fisica", {"Codice Fiscale": "RSSMRA80A01H501U"}, "Rossi Mario"),
            ("b", "giuridica", {"Codice Fiscale": "12345678901"}, "Gianna S.r.l."),
            ("c", "giuridica", {"Codice Fiscale": "10987654321"}, "Bimbo Srl"),
            ("d", "fisica", {"Codice Fiscale": "VRDLRA90A01H501Y"}, "Verdi Laura"),
            ("e", "fisica", {"Codice Fiscale": "VRDLRA91A01H501Z"}, "Verdi Laura"),
        ]
        index = core.build_record_index(candidates)
        self.assertEqual(core.match_checklist_row(index, "mario rossi"), ("ok", ["a"]))
        self.assertEqual(core.match_checklist_row(index, "GIANNA SRL"), ("ok", ["b"]))
        self.assertEqual(core.match_checklist_row(index, "gianna"), ("ok", ["b"]))
        self.assertEqual(core.match_checklist_row(index, "x", "12345678901"), ("ok", ["b"]))
        self.assertEqual(core.match_checklist_row(index, "laura verdi")[0], "ambiguous")
        self.assertEqual(core.match_checklist_row(index, "bimbi srl"), ("similar", ["c"]))
        self.assertEqual(core.match_checklist_row(index, "sconosciuto spa"), ("missing", []))

    def test_matches_person_from_separate_surname_and_name_columns(self):
        headers = ["Cognome", "Nome", "Importo"]
        cells = ["Rossi", "Mario", "100"]
        candidates = [
            ("fisica:cf:rssmra80a01h501u", "fisica",
             {"Codice Fiscale": "RSSMRA80A01H501U"}, "Rossi Mario")
        ]
        index = core.build_record_index(candidates)

        name = core.checklist_person_name(headers, cells)

        self.assertEqual(name, "Rossi Mario")
        self.assertEqual(core.match_checklist_row(index, name), ("ok", [
            "fisica:cf:rssmra80a01h501u"
        ]))

    def test_preserves_single_nominative_column(self):
        self.assertEqual(
            core.checklist_person_name(["Cliente"], ["Mario Rossi"]),
            "Mario Rossi",
        )

    def test_column_mapping_and_values(self):
        headers = [
            "# Mandato", "Cliente", "Capitale Sociale", "% capitale sociale",
            "Versamento in c.to futuro aucap", "TOTALE", "Commissione Apertura",
            "commissione annuale",
        ]
        targets = [
            "COMMISSIONE_APERTURA", "COMMISSIONE_ANNUALE", "CAPITALE SOCIALE",
            "PERCENTUALE CAPITALE SOCIALE", "VERSAMENTO IN C/TO FUTURO AUCAP", "TOTALE",
            "ALTRO",
        ]
        mapping = core.auto_map_columns(targets, headers)
        self.assertEqual(mapping["COMMISSIONE_APERTURA"], 6)
        self.assertEqual(mapping["COMMISSIONE_ANNUALE"], 7)
        self.assertEqual(mapping["CAPITALE SOCIALE"], 2)
        self.assertEqual(mapping["PERCENTUALE CAPITALE SOCIALE"], 3)
        self.assertEqual(mapping["VERSAMENTO IN C/TO FUTURO AUCAP"], 4)
        self.assertEqual(mapping["TOTALE"], 5)
        self.assertNotIn("ALTRO", mapping)
        self.assertEqual(
            core.auto_map_columns(["VERSAMENTO IN CONTO FUTURO AUMENTO DI CAPITALE"], headers),
            {"VERSAMENTO IN CONTO FUTURO AUMENTO DI CAPITALE": 4},
        )
        self.assertEqual(core.checklist_value("PERCENTUALE CAPITALE", "37.5%"), "37.5")
        self.assertEqual(core.checklist_value("QUOTA", "37.5%"), "37.5%")


class MassDraftTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.url = f"sqlite:///{Path(self.temp_dir.name) / 'archive.db'}"

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_save_update_list_delete(self):
        payload = {"phase": 2, "role_data": {"giuridica:1": {"LR": {"role": "Socio"}}}}
        draft_id = core.save_mass_draft(self.url, " Lotto 1 ", payload)
        drafts = core.list_mass_drafts(self.url)
        self.assertEqual([d["name"] for d in drafts], ["Lotto 1"])
        self.assertEqual(drafts[0]["payload"], payload)

        same_id = core.save_mass_draft(self.url, "Lotto 1 bis", {"phase": 3}, draft_id)
        self.assertEqual(same_id, draft_id)
        drafts = core.list_mass_drafts(self.url)
        self.assertEqual(len(drafts), 1)
        self.assertEqual((drafts[0]["name"], drafts[0]["payload"]), ("Lotto 1 bis", {"phase": 3}))

        core.delete_mass_draft(self.url, draft_id)
        self.assertEqual(core.list_mass_drafts(self.url), [])

    def test_name_required(self):
        with self.assertRaises(ValueError):
            core.save_mass_draft(self.url, "  ", {})


if __name__ == "__main__":
    unittest.main()
