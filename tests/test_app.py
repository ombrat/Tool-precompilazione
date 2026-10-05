import io

import pytest
from docx import Document

from precompilazione.app import create_app
from precompilazione.docx_engine import find_fields, fill_document


def make_docx(path):
    d = Document()
    d.add_paragraph("Cliente: {{ragione}} - CF {{cf}}")
    p = d.add_paragraph()
    p.add_run("Premessa: {{pre")
    p.add_run("messa}} fine")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Citta {{citta}}"
    d.sections[0].header.paragraphs[0].text = "H {{ragione}}"
    d.save(path)


def text_of(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    parts += [p.text for p in d.sections[0].header.paragraphs]
    return "\n".join(parts)


def test_find_and_fill(tmp_path):
    src, dst = tmp_path / "a.docx", tmp_path / "b.docx"
    make_docx(src)
    assert find_fields(src) == ["ragione", "cf", "premessa", "citta"]
    fill_document(src, dst, {"ragione": "Acme", "premessa": "uno\ndue", "citta": "Torino"})
    out = text_of(dst)
    assert "Cliente: Acme - CF {{cf}}" in out
    assert "Premessa: uno\ndue fine" in out
    assert "Citta Torino" in out and "H Acme" in out


def test_web_flow(tmp_path):
    src = tmp_path / "a.docx"
    make_docx(src)
    client = create_app(tmp_path / "data").test_client()
    r = client.post("/upload", data={"file": (src.open("rb"), "a.docx"), "name": "Contratto"})
    assert r.status_code == 302
    model_id = r.headers["Location"].split("/")[2]
    r = client.post(f"/modello/{model_id}/configura", data={
        "name": "Contratto", "db_path": "anagrafiche.db", "db_table": "anagrafiche",
        "db_type_column": "tipo_cliente", "db_label_column": "denominazione",
        "source_ragione": "db", "column_ragione": "denominazione",
        "source_cf": "db", "column_cf": "codice_fiscale",
        "source_citta": "db", "column_citta": "citta",
        "source_premessa": "longtext", "default_premessa": "Premessa base"})
    assert r.status_code == 302
    page = client.get(f"/modello/{model_id}/precompila?tipo_cliente=Persona giuridica").data.decode()
    assert "Acme S.r.l." in page and "Mario Rossi" not in page
    rid = page.split('<option value="')[2].split('"')[0]
    r = client.post(f"/modello/{model_id}/precompila", data={
        "tipo_cliente": "Persona giuridica", "anagrafica": rid, "action": "genera",
        "v_premessa": "Premessa custom"})
    assert r.status_code == 200
    out = tmp_path / "o.docx"
    out.write_bytes(r.data)
    text = text_of(out)
    assert "Cliente: Acme S.r.l. - CF 01234567890" in text
    assert "Premessa custom fine" in text and "Torino" in text


def test_upload_rejects_non_docx(tmp_path):
    client = create_app(tmp_path / "data").test_client()
    r = client.post("/upload", data={"file": (io.BytesIO(b"x"), "a.txt")})
    assert r.status_code == 302
    assert client.get("/modello/" + "0" * 32 + "/configura").status_code == 404
