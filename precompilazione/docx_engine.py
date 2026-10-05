"""Rilevamento e sostituzione dei segnaposto {{campo}} nei file Word."""
import re

from docx import Document

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_À-ÿ]+)\s*\}\}")


def _paragraphs(container):
    for p in container.paragraphs:
        yield p
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs(cell)


def iter_paragraphs(doc):
    seen = set()
    parts = [doc]
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            parts.append(part)
    for part in parts:
        for p in _paragraphs(part):
            if id(p._p) not in seen:
                seen.add(id(p._p))
                yield p


def find_fields(path):
    """Restituisce i nomi dei campi nell'ordine di apparizione, senza duplicati."""
    fields = []
    for p in iter_paragraphs(Document(path)):
        for m in PLACEHOLDER.finditer(p.text):
            if m.group(1) not in fields:
                fields.append(m.group(1))
    return fields


def _replace_in_paragraph(p, values):
    # run.text converte "\n" in interruzioni di riga
    full = "".join(r.text for r in p.runs)
    matches = [m for m in PLACEHOLDER.finditer(full) if m.group(1) in values]
    # dal fondo, così le posizioni dei segnaposto precedenti restano valide
    for m in reversed(matches):
        runs = p.runs
        start, end = m.span()
        value = str(values[m.group(1)]).replace("\r\n", "\n").replace("\r", "\n")
        pos = 0
        involved = []
        for run in runs:
            length = len(run.text)
            if length and pos < end and pos + length > start:
                involved.append((run, pos))
            pos += length
        first, first_pos = involved[0]
        last, last_pos = involved[-1]
        before = first.text[: start - first_pos]
        after = last.text[end - last_pos:]
        for run, _ in involved[1:]:
            run.text = ""
        if last is first:
            first.text = before + value + after
        else:
            first.text = before + value
            last.text = after


def fill_document(src, dst, values):
    """Sostituisce i segnaposto con i valori; i campi senza valore restano invariati."""
    doc = Document(src)
    for p in iter_paragraphs(doc):
        _replace_in_paragraph(p, values)
    doc.save(dst)
